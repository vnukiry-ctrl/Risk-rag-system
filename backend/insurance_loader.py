import logging
import os
import json
import re
from datetime import datetime
from typing import Dict, Tuple, Optional
from pypdf import PdfReader
from llm_client import llm_complete, LLM_PROVIDER, MODEL
from dotenv import load_dotenv

load_dotenv()

logger = logging.getLogger(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# Overridable so a single persistent-disk mount (one per service on most
# PaaS platforms) can cover this alongside qdrant_data/log paths, which
# already follow this same os.getenv(..., default) pattern.
DEFAULT_DATA_FOLDER = os.getenv("DATA_FOLDER", os.path.join(BASE_DIR, "data"))

_MULTI_BLANK_RE = re.compile(r"\n{3,}")
_MULTI_SPACE_RE = re.compile(r"[ \t]{2,}")


def _error_record(filename: str, error: str) -> Dict:
    """Build the standard error-metadata shape used whenever a document can't be processed."""
    return {
        "error": error,
        "source_file": filename,
        "extracted_date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def normalize_text(text: str) -> str:
    """Collapse PDF layout whitespace noise without touching real content.

    PyMuPDF preserves the source page layout, so extracted text is full of
    lines that are empty except for spaces (column/table artifacts). Left
    in place, that noise eats into the 8000-char sample sent to the LLM and
    dilutes the chunks used for embedding search.
    """
    # PyMuPDF emits U+FFFD for glyphs it can't map (seen with custom/subset
    # fonts encoding smart quotes and dashes) -- drop rather than keep, since
    # a missing apostrophe ("Lloyds Underwriters") reads better than a
    # visible replacement-character glitch in extracted metadata and answers.
    text = text.replace("�", "")
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = _MULTI_BLANK_RE.sub("\n\n", text)
    text = _MULTI_SPACE_RE.sub(" ", text)
    return text.strip()


# DECISION (DOMAIN-SPECIFIC): deliberately ONLY multi-word structural labels
# ("policy period", "named insured"), not single common words ("insured",
# "policy", "premium"). A first version scored on single words too and it
# backfired: legal wording/insuring-agreement sections use "insured",
# "policy", "insurer" in nearly every sentence, so they out-scored the
# actual declarations page, which states each label once. Multi-word labels
# only appear as page headers/field labels, which is exactly what makes a
# declarations page a declarations page. Rewrite for a new domain (e.g.
# "effective date" / "parties" / "signature block" for contracts).
_DECLARATIONS_KEYWORDS = [
    "named insured", "policy number", "policy no", "policy period",
    "period of insurance", "effective date", "inception date",
    "expiration date", "annual premium", "total premium",
    "limit of liability", "coverage limit", "deductible", "coverholder",
    "declarations", "schedule of",
]

# Declarations pages are dense with actual values (dollar amounts, dates) --
# wording/insuring-agreement prose sections are not, even though they reuse
# "insured"/"policy"/"premium" constantly. These count as strong signals too.
_DATE_PATTERN = re.compile(
    r'\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b'
    r'|\b\d{4}-\d{2}-\d{2}\b'
    r'|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\w*\s+\d{1,2},?\s+\d{4}\b'
    r'|\b\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\w*\s+\d{4}\b',
    re.IGNORECASE,
)


# Grounding/self-verification (gap flagged in review): scalar fields the LLM
# extracts a specific value for, and which we can therefore ask it to back
# with a verbatim quote. Deliberately excludes key_coverages/exclusions
# (lists, not a single span) and notes (a summary, not a quote by design).
_GROUNDED_FIELDS = [
    "policy_number", "insurance_type", "insurance_company", "broker",
    "coverholder", "insured_name", "insured_address", "period_from",
    "period_to", "premium_amount", "coverage_limit", "deductible",
]

_WHITESPACE_RE = re.compile(r"\s+")


def _normalize_for_match(s: str) -> str:
    return _WHITESPACE_RE.sub(" ", s.strip().lower())


_MONTH_NAMES = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def _safe_iso(y: int, mo: int, d: int) -> Optional[str]:
    try:
        return datetime(y, mo, d).strftime("%Y-%m-%d")
    except ValueError:
        return None


def parse_flexible_date(date_str: Optional[str]) -> Optional[str]:
    """Best-effort parse of a free-text policy date into ISO 'YYYY-MM-DD',
    so same-policy documents (renewals/extensions) can be sorted by period
    without comparing raw strings.

    DECISION (DOMAIN-SPECIFIC): ambiguous D/M/Y-vs-M/D/Y slash dates default
    to month-first -- this project's sample declarations pages are US-style
    ("07/01/2026" meaning July 1), falling back to day-first only when the
    first component can't be a month (e.g. "25/03/2026"). Re-check this
    default before reusing on documents from a day-first market. Returns
    None (never a guess) for anything that doesn't match a known shape, so
    "latest version" sorting can fall back to extracted_date instead of
    silently ordering by a misparsed date.
    """
    if not date_str or not isinstance(date_str, str):
        return None
    s = date_str.strip()
    if not s:
        return None

    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})$', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return _safe_iso(y, mo, d)

    # DECISION (real-data gap, found verifying ADR-0012 against actual
    # documents): ordinal day suffixes ("28th July 2026", "1st July 2029")
    # are a real format this document set uses, not a hypothetical -- the
    # optional (?:st|nd|rd|th)? below is what's missing without it: a
    # perfectly parseable date silently fails, sending sort_key to a lower
    # tier (extracted_date) that can rank an actually-later period behind an
    # earlier one that happened to parse.
    m = re.match(
        r'^(\d{1,2})(?:st|nd|rd|th)?[\s-](jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\w*[\s-](\d{4})$',
        s, re.IGNORECASE,
    )
    if m:
        d, mo, y = int(m.group(1)), _MONTH_NAMES[m.group(2).lower()], int(m.group(3))
        return _safe_iso(y, mo, d)

    m = re.match(
        r'^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\w*\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})$',
        s, re.IGNORECASE,
    )
    if m:
        mo, d, y = _MONTH_NAMES[m.group(1).lower()], int(m.group(2)), int(m.group(3))
        return _safe_iso(y, mo, d)

    m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$', s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000 if y < 70 else 1900
        if a > 12:
            day, month = a, b
        else:
            month, day = a, b
        return _safe_iso(y, month, day)

    return None


def verify_grounded_fields(metadata: Dict, quotes: Dict, source_text: str) -> Dict:
    """Verify each extracted field against the quote the LLM claims it came
    from, checking that quote actually appears in the source chunk.

    A field whose quote doesn't check out is NOT trusted silently: it's
    nulled in `metadata` in place, and the failure (with the original value,
    for debugging) is recorded in the returned per-field verification map.
    """
    normalized_source = _normalize_for_match(source_text)
    verification: Dict[str, Dict] = {}

    for field in _GROUNDED_FIELDS:
        value = metadata.get(field)
        if value is None:
            verification[field] = {"status": "not_applicable"}
            continue

        quote = quotes.get(field) if isinstance(quotes, dict) else None
        if not quote or not isinstance(quote, str) or not quote.strip():
            verification[field] = {"status": "unverified_missing_quote", "original_value": value}
            metadata[field] = None
            continue

        normalized_quote = _normalize_for_match(quote)
        if normalized_quote in normalized_source:
            verification[field] = {"status": "verified", "quote": quote}
        else:
            verification[field] = {
                "status": "unverified_quote_not_found",
                "quote": quote,
                "original_value": value,
            }
            metadata[field] = None

    return verification


_LABEL_RE = re.compile(
    r'(?:' + '|'.join(re.escape(kw) for kw in _DECLARATIONS_KEYWORDS) + r')\s*:',
    re.IGNORECASE,
)


def _score_window(window: str) -> int:
    # DECISION: a bare keyword mention isn't enough -- "as set forth in the
    # Declarations" or "the Effective Date of this Policy" reference these
    # terms constantly throughout ordinary wording/insuring-agreement prose,
    # not just on the actual declarations page (confirmed: those bare
    # mentions alone outscored the real page on the 269K-char D&O EPL
    # policy). What a declarations page actually looks like is a LABEL
    # immediately followed by a colon and a value ("Named Insured: The
    # Board of..."), so only that pattern counts -- a structural test, not
    # one tuned to any specific document. $/date hits still only count
    # alongside at least one such label, for the same reason as before: a
    # bare cluster of $ signs (a coverage-limits schedule) isn't evidence
    # of a declarations page on its own.
    label_hits = len(_LABEL_RE.findall(window))
    if label_hits == 0:
        return 0
    dollar_hits = window.count('$')
    date_hits = len(_DATE_PATTERN.findall(window))
    return 5 * label_hits + 2 * dollar_hits + 2 * date_hits


def select_declarations_window(text: str, max_chars: int = 8000, window_size: int = 2000, window_overlap: int = 200) -> str:
    """Classify-then-target: instead of blindly taking text[:max_chars] (which
    can land entirely on front-matter boilerplate for a document whose real
    declarations page sits tens of thousands of characters in), score
    overlapping windows by declarations-vocabulary density and keep only the
    highest-scoring ones, up to max_chars total.

    Selected windows are restored to their original document order before
    being joined, so the LLM still reads a coherent passage rather than
    disconnected fragments in an arbitrary order.
    """
    if len(text) <= max_chars:
        return text

    step = window_size - window_overlap
    windows = []
    start = 0
    while start < len(text):
        end = min(start + window_size, len(text))
        windows.append((start, end, text[start:end]))
        if end == len(text):
            break
        start += step

    scored = sorted(windows, key=lambda w: _score_window(w[2]), reverse=True)

    selected = []
    total = 0
    for w in scored:
        if selected and total + len(w[2]) > max_chars:
            continue
        selected.append(w)
        total += len(w[2])
        if total >= max_chars:
            break

    selected.sort(key=lambda w: w[0])
    logger.debug(
        "select_declarations_window: kept %d/%d windows (%d chars) from a %d-char document",
        len(selected), len(windows), total, len(text),
    )
    return "\n\n".join(w[2] for w in selected)


class InsuranceDocumentParserLLM:
    def extract_metadata_with_llm(self, text: str, filename: str, extraction_strategy: str = "head_truncate") -> Dict:
        """Use Groq to intelligently extract metadata from insurance document"""

        if extraction_strategy == "classify_then_target":
            text_sample = select_declarations_window(text)
        else:
            # DECISION (UNIVERSAL, value needs re-checking per project): hard
            # truncation before the fields-of-interest are guaranteed to appear.
            # A doc whose key fields sit past char 8000 (long riders/addenda)
            # silently loses them here -- verify against your actual document
            # lengths rather than reusing 8000.
            text_sample = text[:8000]

        # DECISION (DOMAIN-SPECIFIC): entire prompt below -- field list,
        # few-shot hints (policy-number formats, insurer/broker phrasing) --
        # is insurance vocabulary. Rewrite completely for a new domain; the
        # pattern worth keeping is "explicit field list + null-if-missing +
        # return-only-JSON", not the specific fields or hints.
        prompt = f"""You are an insurance document expert. Extract ONLY valid JSON from this insurance document. Return the JSON object only, no other text.

DOCUMENT FILENAME: {filename}
DOCUMENT TEXT:
{text_sample}

Extract these fields. Return ONLY this JSON structure (no markdown, no explanation):
{{
    "policy_number": "Policy number/ID (look for 'Policy No', 'Policy #', 'ALCOG', 'BW', 'XLKR' formats) or null",
    "insurance_type": "Type of insurance (e.g., Personal Accident, General Liability, Kidnap & Ransom) or null",
    "insurance_company": "Name(s) of the INSURER/Insurance Company/Underwriter/Syndicate. If MULTIPLE insurers/underwriters participate (co-insurance, subscription market, syndicates), list ALL of them, comma-separated, not just the first -- or null",
    "broker": "Name of BROKER/Agent who arranged the policy or null",
    "coverholder": "Name of COVERHOLDER (may be same as broker) or null",
    "insured_name": "Name of policyholder/insured party or null",
    "insured_address": "Address of insured or null",
    "period_from": "Start date of policy or null",
    "period_to": "End/expiration date of policy or null",
    "document_role": "One of: original, renewal, endorsement, extension, cancellation -- based on how the document itself describes its relationship to the policy (e.g. 'Renewal Certificate', 'Endorsement No.', 'Extension of Cover'). If there's no such indicator, use 'original'. Never guess between renewal/endorsement/extension without wording support -- use null instead",
    "premium_amount": "Annual premium amount (numbers only, no currency) or null",
    "coverage_limit": "Main coverage limit amount or null",
    "deductible": "Deductible amount or null",
    "key_coverages": ["coverage1", "coverage2", "coverage3"],
    "exclusions": ["exclusion1", "exclusion2"],
    "notes": "Any other important details or null",
    "field_quotes": {{
        "policy_number": "verbatim text from DOCUMENT TEXT that this value came from, or null",
        "insurance_type": "verbatim quote or null",
        "insurance_company": "verbatim quote or null",
        "broker": "verbatim quote or null",
        "coverholder": "verbatim quote or null",
        "insured_name": "verbatim quote or null",
        "insured_address": "verbatim quote or null",
        "period_from": "verbatim quote or null",
        "period_to": "verbatim quote or null",
        "premium_amount": "verbatim quote or null",
        "coverage_limit": "verbatim quote or null",
        "deductible": "verbatim quote or null"
    }}
}}

CRITICAL RULES:
- Return ONLY the JSON object, absolutely nothing else
- Use null for missing fields (not "N/A" or "Unknown")
- For POLICY NUMBER: Look for "Policy No", "Policy #", "ALCOG", "BW", "XLKR", "Assigned:", policy numbers in boxes/headers
- For INSURANCE COMPANY: Look for "Insurance Company", "Underwriter", "Insured with", "Effected with", "Lloyd's", "AXA", "Aviva", "Allianz", or similar -- if the document names several participating insurers/underwriters (co-insurance, subscription market), include ALL of them, not just the first one found
- For INSURANCE COMPANY's field_quotes entry specifically: if multiple insurers are listed, the quote must be the verbatim source passage that contains all of their names (copy it as it appears, spanning multiple lines/commas if needed) -- not a quote for only the first insurer
- For BROKER: Look for "Broker", "Agent", "Through", "Arranged by", "Via", or company names that arrange insurance
- For COVERHOLDER: Look for "Lloyd's Approved Coverholder", "Coverholder", "Administrator"
- For DOCUMENT_ROLE: only use "renewal"/"endorsement"/"extension"/"cancellation" when the document explicitly labels itself that way (title, header, or a sentence like "This endorsement amends..."); otherwise use "original"
- Be accurate - only extract text that actually appears in the document
- Return valid JSON that can be parsed
- For amounts, return only numbers (e.g., "25000" not "$25,000")
- For every field listed in "field_quotes": copy the exact text VERBATIM from DOCUMENT TEXT above (same spelling, spacing, punctuation) that justifies that field's value -- it will be checked programmatically against the source text. If you cannot find a verbatim supporting quote for a field, set BOTH that field's value and its field_quotes entry to null rather than guessing."""

        try:
            response_text = llm_complete(
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
                max_tokens=2500,
            ).strip()
            
            # Remove markdown if present
            response_text = re.sub(r'```json\n?', '', response_text)
            response_text = re.sub(r'```\n?', '', response_text)
            response_text = response_text.strip()
            
            # Validate JSON before parsing
            if not response_text.startswith('{'):
                logger.warning("%s: LLM response doesn't start with JSON", filename)
                return _error_record(filename, "Invalid LLM response format")
            
            # Parse JSON
            metadata = json.loads(response_text)

            # Pull the per-field quotes out before anything else touches
            # metadata -- they're grounding evidence, not an extracted field.
            field_quotes = metadata.pop("field_quotes", None) or {}

            # Ensure all expected fields exist
            default_fields = {
                "policy_number": None,
                "insurance_type": None,
                "insurance_company": None,
                "broker": None,
                "coverholder": None,
                "insured_name": None,
                "insured_address": None,
                "period_from": None,
                "period_to": None,
                "document_role": "original",
                "premium_amount": None,
                "coverage_limit": None,
                "deductible": None,
                "key_coverages": [],
                "exclusions": [],
                "notes": None
            }
            
            # Merge with defaults
            for key in default_fields:
                if key not in metadata:
                    metadata[key] = default_fields[key]

            # Grounding pass: null out (don't silently trust) any field whose
            # claimed source quote doesn't actually appear in the text the
            # LLM was given.
            field_verification = verify_grounded_fields(metadata, field_quotes, text_sample)
            flagged_fields = [
                field for field, v in field_verification.items()
                if v["status"] not in ("verified", "not_applicable")
            ]
            if flagged_fields:
                logger.warning("%s: unverified fields nulled: %s", filename, flagged_fields)

            # Parsed after grounding, using the (possibly nulled) final
            # values -- an unverified period_from/period_to shouldn't get a
            # confidently-parsed ISO date sitting next to a null raw value.
            metadata["period_from_iso"] = parse_flexible_date(metadata.get("period_from"))
            metadata["period_to_iso"] = parse_flexible_date(metadata.get("period_to"))

            # Add metadata
            metadata["extracted_date"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            metadata["source_file"] = filename
            metadata["extraction_method"] = f"{LLM_PROVIDER} LLM ({MODEL})"
            metadata["extraction_strategy"] = extraction_strategy
            metadata["field_verification"] = field_verification
            metadata["flagged_fields"] = flagged_fields

            return metadata

        except json.JSONDecodeError as e:
            logger.error("%s: JSON parse error: %s | response was: %s", filename, e, response_text[:200])
            return _error_record(filename, f"JSON parse error: {str(e)}")
        except Exception as e:
            logger.error("%s: LLM API call failed", filename, exc_info=True)
            return _error_record(filename, str(e))

    def parse_document(self, text: str, filename: str, extraction_strategy: str = "head_truncate") -> Tuple[Dict, str]:
        """Parse insurance document using Groq LLM"""
        metadata = self.extract_metadata_with_llm(text, filename, extraction_strategy=extraction_strategy)
        return metadata, text


def extract_text_from_pdf(pdf_path: str) -> str:
    """Extract text from PDF using PyMuPDF (better CID font handling)"""
    try:
        import fitz

        text = ""
        doc = fitz.open(pdf_path)
        num_pages = len(doc)
        logger.debug("Reading %d pages from %s", num_pages, pdf_path)

        for page_num in range(num_pages):
            page = doc[page_num]
            extracted = page.get_text()
            if extracted:
                text += extracted

        doc.close()
        return text
    except Exception as e:
        logger.error("Failed to read PDF %s", pdf_path, exc_info=True)
        return ""


def extract_text_from_docx(docx_path: str) -> str:
    """Extract text from a DOCX file (python-docx) -- paragraphs only.

    DECISION (matches 5.1's ladder in docs/learning-guide.md): tables and
    embedded objects aren't walked here. DOCX is structured XML, so unlike
    the PDF path this can be extended precisely (python-docx exposes
    doc.tables) the moment a real document needs it -- no guessing required,
    just not built until a document actually has table content worth the
    coverage.
    """
    try:
        from docx import Document

        doc = Document(docx_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error("Failed to read DOCX %s", docx_path, exc_info=True)
        return ""


def extract_text_from_image(image_path: str) -> str:
    """Extract text from a scanned image/photo via Tesseract OCR.

    DECISION (docs/learning-guide.md 5.1): this is the actual "no text layer
    at all" case the OCR row exists for -- a standalone image has nothing
    else to fall back to, unlike a PDF with a real text layer.

    DECISION (docs/learning-guide.md 5.6, "fail loudly not silently"): the
    tesseract binary is a separate OS-level install, not something `pip
    install pytesseract` provides -- pytesseract is only a wrapper around
    it. A missing binary is an environment problem, not "this one image had
    no readable text," so pytesseract.TesseractNotFoundError is deliberately
    NOT caught here. It propagates to load_insurance_documents' existing
    per-document try/except (§1.5), which records it with its own real,
    actionable message ("tesseract is not installed or it's not in your
    PATH") instead of being folded into the generic empty-text case below.
    Any other failure (corrupt/unreadable image) IS caught here, since that
    really is specific to this one file.
    """
    from PIL import Image
    import pytesseract

    try:
        image = Image.open(image_path)
        return pytesseract.image_to_string(image)
    except pytesseract.TesseractNotFoundError:
        raise
    except Exception as e:
        logger.error("Failed to OCR image %s", image_path, exc_info=True)
        return ""


# Format-dispatch layer (docs/learning-guide.md 5.1): route by extension to
# the right extractor, normalize every format to the same plain-text shape
# before it reaches normalize_text()/chunking/metadata extraction below, so
# nothing downstream needs to know or care what format a document arrived in.
_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp"}
_TEXT_EXTRACTORS = {
    ".pdf": extract_text_from_pdf,
    ".docx": extract_text_from_docx,
    **{ext: extract_text_from_image for ext in _IMAGE_EXTENSIONS},
}
_PLAIN_TEXT_EXTENSIONS = {".txt"}


def _is_supported_extension(ext: str) -> bool:
    return ext in _PLAIN_TEXT_EXTENSIONS or ext in _TEXT_EXTRACTORS


def load_insurance_documents(data_folder: Optional[str] = None, extraction_strategy: str = "head_truncate") -> Dict:
    """Load and parse all insurance documents using Groq LLM."""
    if data_folder is None:
        data_folder = DEFAULT_DATA_FOLDER
    elif not os.path.isabs(data_folder):
        data_folder = os.path.join(BASE_DIR, data_folder)

    parser = InsuranceDocumentParserLLM()
    results = {
        "metadata": [],
        "parent_chunks": []
    }

    if not os.path.exists(data_folder):
        logger.warning("Data folder not found: %s", data_folder)
        return results

    files = [f for f in os.listdir(data_folder) if os.path.isfile(os.path.join(data_folder, f))]

    if not files:
        logger.warning("No files found in %s", data_folder)
        return results

    logger.info("Parsing %d documents from %s", len(files), data_folder)

    for filename in files:
        filepath = os.path.join(data_folder, filename)
        ext = os.path.splitext(filename)[1].lower()

        try:
            if not _is_supported_extension(ext):
                # Fail loudly, not silently (docs/learning-guide.md 5.6): the
                # old filter just excluded anything that wasn't .txt/.pdf
                # from `files` above, so an unrecognized format vanished with
                # no record anywhere. This is layer 1 (structured logging) of
                # that pattern -- a queryable rate (layer 2) is deferred to
                # Milestone 7, same as the README roadmap notes.
                logger.warning("%s: unsupported file format (%s)", filename, ext or "no extension")
                results["metadata"].append(
                    _error_record(filename, f"Unsupported file format: {ext or '(no extension)'}")
                )
                continue

            if ext in _PLAIN_TEXT_EXTENSIONS:
                with open(filepath, 'r', encoding='utf-8') as f:
                    text = f.read()
            else:
                text = _TEXT_EXTRACTORS[ext](filepath)
                if not text.strip():
                    logger.warning("%s: no text could be extracted (%s)", filename, ext)
                    results["metadata"].append(
                        _error_record(filename, f"No text could be extracted from {ext} file")
                    )
                    continue

            text = normalize_text(text)

            metadata, parent = parser.parse_document(text, filename, extraction_strategy=extraction_strategy)

            results["metadata"].append(metadata)

            results["parent_chunks"].append({
                "content": parent,
                "metadata": metadata,
                "source_file": filename,
                "type": "parent"
            })

            if "error" in metadata:
                logger.warning("%s: metadata extraction failed: %s", filename, metadata["error"])
            else:
                logger.info(
                    "Extracted %s: policy=%s type=%s insurer=%s",
                    filename,
                    metadata.get('policy_number', 'N/A'),
                    metadata.get('insurance_type', 'N/A'),
                    metadata.get('insurance_company', 'N/A'),
                )

        except Exception as e:
            logger.error("Failed to parse %s", filename, exc_info=True)
            results["metadata"].append(_error_record(filename, str(e)))

    return results


if __name__ == "__main__":
    print("=" * 70)
    print("INSURANCE DOCUMENT PARSER - GROQ LLM VERSION")
    print("Using Groq for intelligent metadata extraction")
    print("=" * 70)
    
    results = load_insurance_documents()
    
    print("\n" + "=" * 70)
    print("SUMMARY")
    print("=" * 70)
    
    print(f"\nPolicies Loaded: {len(results['metadata'])}")
    print(f"Parent Chunks: {len(results['parent_chunks'])}")
    
    if len(results['metadata']) > 0:
        print("\n" + "=" * 70)
        print("EXTRACTED METADATA")
        print("=" * 70)
        
        for meta in results['metadata']:
            print(f"\n{'='*70}")
            print(f"File: {meta.get('source_file', 'Unknown')}")
            print(f"{'='*70}")
            print(f"  Policy #: {meta.get('policy_number', 'N/A')}")
            print(f"  Type: {meta.get('insurance_type', 'N/A')}")
            print(f"  Insurer: {meta.get('insurance_company', 'N/A')}")
            print(f"  Broker: {meta.get('broker', 'N/A')}")
            print(f"  Coverholder: {meta.get('coverholder', 'N/A')}")
            print(f"  Insured: {meta.get('insured_name', 'N/A')}")
            print(f"  Address: {meta.get('insured_address', 'N/A')}")
            print(f"  Period: {meta.get('period_from', 'N/A')} to {meta.get('period_to', 'N/A')}")
            print(f"  Premium: ${meta.get('premium_amount', 'N/A')}")
            print(f"  Coverage Limit: {meta.get('coverage_limit', 'N/A')}")
            print(f"  Deductible: {meta.get('deductible', 'N/A')}")
            if meta.get('key_coverages'):
                print(f"  Coverages: {', '.join(meta.get('key_coverages', []))}")
            if meta.get('exclusions'):
                print(f"  Exclusions: {', '.join(meta.get('exclusions', []))}")
            print(f"  Method: {meta.get('extraction_method', 'N/A')}")
    
    print("\n" + "=" * 70)