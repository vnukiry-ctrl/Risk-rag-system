"""Unit tests for the Milestone 5.1 multi-format dispatch layer.

Unlike test_quality.py (which talks to the live API -- see conftest.py's
require_running_backend), these run directly against insurance_loader.py:
extraction and format-dispatch are pure functions, so there's no reason to
pay for a live backend + LLM call just to check that a DOCX paragraph comes
back as text, or that an unrecognized extension gets logged instead of
silently dropped.
"""
import os
import shutil
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from PIL import Image, ImageDraw

from insurance_loader import (
    extract_text_from_docx,
    extract_text_from_image,
    _is_supported_extension,
    load_insurance_documents,
)

TESSERACT_INSTALLED = shutil.which("tesseract") is not None


def test_docx_extraction_returns_paragraph_text(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Policy Number: BW240599")
    doc.add_paragraph("Insured: Mount Royal University")
    doc.save(str(docx_path))

    text = extract_text_from_docx(str(docx_path))

    assert "Policy Number: BW240599" in text
    assert "Mount Royal University" in text


def test_docx_extraction_returns_empty_string_on_bad_file(tmp_path):
    bad_path = tmp_path / "not_really.docx"
    bad_path.write_text("this is not a real docx file")

    assert extract_text_from_docx(str(bad_path)) == ""


def test_supported_extensions():
    assert _is_supported_extension(".txt")
    assert _is_supported_extension(".pdf")
    assert _is_supported_extension(".docx")
    assert _is_supported_extension(".jpg")
    assert _is_supported_extension(".png")
    assert not _is_supported_extension(".xyz")
    assert not _is_supported_extension("")


def _make_text_image(path, text):
    image = Image.new("RGB", (400, 100), color="white")
    draw = ImageDraw.Draw(image)
    draw.text((10, 40), text, fill="black")
    image.save(path)


@pytest.mark.skipif(
    not TESSERACT_INSTALLED,
    reason="Tesseract OCR binary not installed on this host -- pip installs "
    "pytesseract (the wrapper) but not the engine itself. See "
    "https://github.com/tesseract-ocr/tesseract#installing-tesseract. "
    "This is the exact gap docs/learning-guide.md 5.6 calls 'fail loudly, "
    "not silently' -- skipped here, not silently passed.",
)
def test_image_ocr_extracts_text(tmp_path):
    image_path = tmp_path / "scan.png"
    _make_text_image(str(image_path), "POLICY BW240599")

    text = extract_text_from_image(str(image_path))

    assert "BW240599" in text


def test_missing_tesseract_binary_fails_loudly_not_silently(tmp_path):
    """Documents this environment's real current state: the OCR *code* is
    correct and wired in, but the Tesseract *binary* isn't installed here.
    Per docs/learning-guide.md 5.6, that must surface as a distinct,
    actionable error record -- not an empty "no text found" result that
    looks the same as a genuinely blank scan.
    """
    if TESSERACT_INSTALLED:
        pytest.skip("Tesseract IS installed on this host -- this test only "
                    "verifies the missing-binary failure path.")

    image_path = tmp_path / "scan.png"
    _make_text_image(str(image_path), "POLICY BW240599")

    results = load_insurance_documents(data_folder=str(tmp_path))

    assert len(results["metadata"]) == 1
    record = results["metadata"][0]
    assert record["source_file"] == "scan.png"
    assert "tesseract" in record["error"].lower()


def test_unsupported_format_is_logged_not_silently_dropped(tmp_path):
    (tmp_path / "notes.xyz").write_text("some content in an unsupported format")

    results = load_insurance_documents(data_folder=str(tmp_path))

    assert len(results["metadata"]) == 1
    record = results["metadata"][0]
    assert record["source_file"] == "notes.xyz"
    assert "Unsupported file format" in record["error"]
    assert len(results["parent_chunks"]) == 0
